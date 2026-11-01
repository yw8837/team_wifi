"""Offline public-data experiments. Never joins public identities to app users."""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime, timezone
import hashlib
import itertools
import json
from pathlib import Path
import platform
import zipfile
import xml.etree.ElementTree as ET

import joblib
import numpy as np
import pandas as pd
import sklearn
from sklearn.dummy import DummyClassifier
from sklearn.impute import SimpleImputer
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import average_precision_score, brier_score_loss, log_loss, roc_auc_score
from sklearn.model_selection import GroupShuffleSplit
from sklearn.pipeline import make_pipeline
from sklearn.preprocessing import StandardScaler
from lightgbm import LGBMClassifier

HERE = Path(__file__).resolve().parent
FEATURES = ["Administrative", "Administrative_Duration", "Informational", "Informational_Duration",
            "ProductRelated", "ProductRelated_Duration", "SpecialDay", "Weekend"]
SEED = 20261101
BUSINESS = ["InvoiceNo", "StockCode", "Description", "Quantity", "InvoiceDate", "UnitPrice", "CustomerID", "Country"]
LIMITATIONS = [
    "UCI 과거 세션의 구매 여부 분류 실험이며 한국 육아용품 구매 확률이 아닙니다.",
    "방문 수와 체류시간은 세션 종료 집계입니다. 진행 중 세션의 실시간 성능은 검증하지 않았습니다.",
    "확률의 국내 서비스 보정과 외부 검증이 없으며 할인·이탈·공구 성사 원인을 설명하지 않습니다.",
    "PageValues, BounceRates, ExitRates, Revenue는 입력에서 제외했습니다.",
]


def sha256(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


def save_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")


def canonical(frame):
    """Normalize 8 business fields, including multiplicity, never the export index."""
    f = frame.rename(columns={"Invoice": "InvoiceNo", "Price": "UnitPrice", "Customer ID": "CustomerID"}).copy()
    for c in ["InvoiceNo", "StockCode", "CustomerID"]:
        f[c] = f[c].fillna("").astype(str).str.replace(r"\.0$", "", regex=True)
    for c in ["Description", "Country"]:
        f[c] = f[c].fillna("").astype(str)
    f["InvoiceDate"] = pd.to_datetime(f["InvoiceDate"], format="mixed").dt.strftime("%Y-%m-%d %H:%M:%S")
    for c in ["Quantity", "UnitPrice"]:
        f[c] = pd.to_numeric(f[c]).map(lambda n: format(n, ".12g"))
    # Tuples preserve exact canonical values; no hash collision assumption in overlap audit.
    return Counter(f[BUSINESS].itertuples(index=False, name=None))


def material_inventory(data_root, out):
    export = data_root / "notion_export"
    inventory = []
    original_notes = []
    requirements = []
    for p in sorted(export.rglob("*")):
        if not p.is_file():
            continue
        name = p.name
        status, reason = "reference", "보존된 출처·복원·분석 자료이며 새 원본 분석과 대조; 제품 실행 코드로 자동 채택하지 않음"
        content = ""
        if p.suffix.lower() in {".md", ".txt", ".json", ".csv", ".py", ".ipynb"}:
            content = p.read_text(encoding="utf-8-sig", errors="replace")
        if "Elice" in name or "elice" in name.lower():
            status, reason = "disabled", "사용자 지시에 따라 호출·복구·연결 점검 금지"
        elif "기획안_v1" in name or "기획안_v2" in name:
            status, reason = "superseded", "v3 및 현재 모두의 육아 범위 이전 기획; 원본 보존, 재생성하지 않음"
        elif "기획안_v3" in name or name.startswith("요구사항 정의서 상세"):
            status, reason = "implementation_reference", "관심사·탐색·구매/참여·권한·신뢰·추천·GEO 요구를 현재 기능 대조표에 사용; 모든 과거 기능의 출시를 뜻하지 않음"
        elif name.endswith(".docx"):
            status, reason = "implementation_reference", "기업 원문 공동구매 주제와 활용 방안·라이선스 사전 확인 요구를 4종 분석/추론 기능으로 구체화"
        elif name.startswith(("D1", "D2", "D3", "D4", "source_", "column_", "download_")):
            status, reason = "experiment", "네 데이터의 출처·스키마·위험 검토; 실제 계산은 raw/processed 원본으로 재실행"
        elif name.endswith(".zip") or name.startswith("normalized_part"):
            status, reason = "reference", "복원/분석 패키지: 경로·크기·해시와 ZIP 구성 조사, 이미 확보된 원본을 중복 재다운로드하지 않음"
        elif "사용자 " in name or "관리자 " in name:
            status, reason = "implementation_reference", "요구사항 개별 페이지 읽기; 구체 기능의 구현/시뮬레이션/보류 상태는 메인 기능 대조표로 확인"
        if p.suffix.lower() == ".docx":
            with zipfile.ZipFile(p) as z:
                root = ET.fromstring(z.read("word/document.xml"))
                ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                paras = ["".join(t.text or "" for t in para.findall(".//w:t", ns)) for para in root.findall(".//w:p", ns)]
                relations = ET.fromstring(z.read("word/_rels/document.xml.rels"))
                links = sorted({r.attrib["Target"] for r in relations if r.attrib.get("TargetMode") == "External" and r.attrib.get("Target", "").startswith("http")})
                original_notes.append({"source": name, "sha256": sha256(p), "paragraphs": [v for v in paras if v.strip()], "external_links": links})
        if p.suffix.lower() == ".csv" and name.startswith("요구사항 정의서 상세") and "_all" in name:
            rf = pd.read_csv(p, keep_default_na=False, encoding="utf-8-sig")
            requirements = [{"feature": r["기능명"], "description": r["설명"], "priority": r["중요도"]} for _, r in rf.iterrows()]
        record = {"path": str(p.relative_to(data_root)), "bytes": p.stat().st_size, "sha256": sha256(p), "status": status, "reason": reason}
        if content:
            record["read_scope"] = "whole UTF-8 text read; original hashed"
            record["characters_read"] = len(content)
            if "Elice" in content or "elice" in content.lower():
                record["disabled_reference"] = "본문에 남은 Elice 지시는 실행하지 않음"
            if p.suffix.lower() == ".ipynb":
                notebook = json.loads(content)
                record["notebook_read"] = {"cells": len(notebook.get("cells", [])), "stored_outputs": sum(len(c.get("outputs", [])) for c in notebook.get("cells", [])), "execution": "stored sources/outputs read only; not re-executed as new evidence"}
        if p.suffix.lower() == ".xlsx":
            import openpyxl
            workbook = openpyxl.load_workbook(p, read_only=True, data_only=True)
            record["workbook_read"] = {s.title: {"rows": s.max_row, "columns": s.max_column, "header": [str(v) if v is not None else "" for v in next(s.values)]} for s in workbook}
            workbook.close()
        if p.suffix.lower() == ".zip":
            try:
                with zipfile.ZipFile(p) as z:
                    record["archive_members"] = [{"path": i.filename, "bytes": i.file_size} for i in z.infolist()]
            except zipfile.BadZipFile:
                record["read_scope"] = "multipart archive fragment, inventoried only; restored source already available"
        inventory.append(record)
    save_json(out / "materials_inventory.json", {"read_at": "2026-11-01", "scope": "local versioned Notion export; not a claim of live Notion/blog access", "files": inventory})
    save_json(out / "company_source_extract.json", original_notes)
    save_json(out / "requirements_source.json", requirements)
    source_matrix = []
    for i, r in enumerate(requirements, 1):
        feature = r["feature"]
        if feature in {"상품 추천 기준모델 및 비교 평가", "공개데이터 활용 범위 구분"}:
            stance = "현재실험"
            reason = "D1-D4 분석/누수검사와 시간순 묶음완성 평가를 실제 실행; 과거 영국 상품 실험을 현재 육아 상품과 분리"
            evidence = "summary.json; evaluation.json; overlap_audit.json; bundle_rules.json; verify-run.log"
        elif any(k in feature for k in ["정산", "분담", "상환", "환불", "납부", "지급", "결제", "배송", "수령", "공급", "재고", "예약", "보호", "구매확정", "선지출", "심사", "신뢰 지표"]):
            stance = "후속검증"
            reason = "이 요구의 이행·금융·권한·공급 정답은 D1-D4에 없음; 실제 앱 코드/시험 및 필요한 사업자 연동을 별도로 대조"
            evidence = "requirements_source.json (original requirement); external integration not verified by this pipeline"
        else:
            stance = "참조"
            reason = "원문 요구를 관심사·조건 우선 탐색·실제 상품 근거·권한·GEO 설계 기준으로 사용; 현재 기능 완료는 메인 코드/화면 시험으로 판정"
            evidence = "requirements_source.json; IMPLEMENTATION.md"
        source_matrix.append({"source_row": i + 1, **r, "source_stance": stance, "reason": reason, "evidence": evidence, "ui_implementation_status": "미판정 (main-owned app verification)"})
    save_json(out / "requirements_traceability.json", source_matrix)
    pd.DataFrame(source_matrix).to_csv(out / "requirements_traceability.csv", index=False, encoding="utf-8-sig")
    lines = ["# 원문 요구사항 101개와 데이터 작업의 연결", "", "기준일 2026-11-01. 이 표의 현재실험·참조·후속검증은 데이터 파이프라인의 근거 범위다. **101개 모두 UI 구현 상태는 미판정**이며 구현 완료/제외 승인을 뜻하지 않는다. 메인이 실제 코드·화면·시험으로 별도 대조한다. 원문 설명 전체와 행번호는 같은 폴더 requirements_traceability.csv/json 및 requirements_source.json에 보존한다.", "", "| 원문 행 | 기능 | 데이터 활용 상태 | 현재 근거 및 추가 확인 |", "|---:|---|---|---|"]
    for r in source_matrix:
        lines.append(f"| {r['source_row']} | {r['feature']} | {r['source_stance']} | {r['reason']} |")
    (out / "requirements_traceability.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    return {"files_inventoried": len(inventory), "requirements_read": len(requirements), "company_documents_read": len(original_notes)}


def consumer_analysis(path):
    f = pd.read_csv(path, keep_default_na=False)
    category = f["Purchase_Category"].value_counts()
    grouped = f.groupby(["Discount_Sensitivity", "Purchase_Channel"], dropna=False).size()
    return {"rows": len(f), "columns": len(f.columns), "categories": [{"category": k, "customers": int(v)} for k, v in category.items()],
            "discount_sensitivity": {k: int(v) for k, v in f["Discount_Sensitivity"].value_counts().items()},
            "discount_channel": [{"sensitivity": k[0], "channel": k[1], "customers": int(v)} for k, v in grouped.items()],
            "none_categories_preserved": {c: int(f[c].eq("None").sum()) for c in ["Social_Media_Influence", "Engagement_with_Ads"]},
            "application": "선택 관심사·쿠폰/혜택 탐색 UI의 세분화 근거. 할인 민감도는 추정 가입 속성이 아니라 명시 선택이며 할인 효과를 주장하지 않음.",
            "limitations": ["수집·표집·실측/합성 여부 미확인", "Occupation이 High/Middle이므로 직업 타기팅에서 제외", "Purchase_Intent는 구매 여부가 아닌 명목 범주", "국내 육아 시장으로 일반화 불가"]}


def demand_backtest(monthly):
    """Two baselines on 11 held-out full months; no future orders used."""
    full = monthly[~monthly.partial_month].reset_index(drop=True)
    observations = []
    for i in range(13, len(full)):
        current = full.iloc[i]
        observations.append({"month": current["month"], "actual_units": int(current["units"]),
                             "previous_month_units": int(full.iloc[i - 1]["units"]),
                             "seasonal_12month_units": int(full.iloc[i - 12]["units"])})
    result = {}
    actual = np.array([v["actual_units"] for v in observations], dtype=float)
    for model in ["previous_month_units", "seasonal_12month_units"]:
        predicted = np.array([v[model] for v in observations], dtype=float)
        error = np.abs(actual - predicted)
        result[model] = {"mae_units": float(error.mean()), "wape": float(error.sum() / actual.sum())}
    return {"target": "monthly positive sold units, not latent market demand", "test_months": len(observations),
            "evaluation": "expanding rolling origin; each prediction uses only prior months; no parameter fitting or test-based model selection",
            "observations": observations, "models": result,
            "limitations": ["24 full months only", "no inventory/stockout/exposure/price covariates", "aggregate UK gifts/wholesale", "2011-12 excluded as partial month", "no current childcare demand forecast"]}


def retail_analysis(data_root, out):
    normalized = data_root / "processed/retail_ii_normalized.csv.gz"
    restore = read_json(data_root / "notion_export/restore_manifest.json")
    expected = next(v["output_sha256"] for v in restore["datasets"] if v["id"] == "normalized")
    assert sha256(normalized) == expected, "Normalized checksum mismatch; regenerate from XLSX before use"
    f = pd.read_csv(normalized, dtype={"InvoiceNo": str, "StockCode": str, "CustomerID": str}, keep_default_na=False, encoding="utf-8-sig")
    assert len(f) == 1067371
    sheet_counts = f.groupby("source_sheet").size().to_dict()
    assert sheet_counts == {"Year 2009-2010": 525461, "Year 2010-2011": 541910}
    assert all(g["source_row"].tolist() == list(range(2, len(g) + 2)) for _, g in f.groupby("source_sheet", sort=False))
    # Recompute all flags from the eight business fields rather than trust old flags.
    f["InvoiceDate"] = pd.to_datetime(f["InvoiceDate"])
    duplicated = f.duplicated(BUSINESS)
    product = (f["Quantity"].gt(0) & f["UnitPrice"].gt(0) & ~f["InvoiceNo"].str.startswith("C") &
               f["Description"].ne("") & f["StockCode"].str.fullmatch(r"\d{5}[A-Za-z]{0,2}") & ~duplicated)
    assert int(product.sum()) == 1003213
    d2 = pd.read_csv(data_root / "raw/02_kaggle_retail/online_retail.csv", dtype=str, keep_default_na=False)
    a = canonical(d2)
    b = canonical(f[f.source_sheet.eq("Year 2010-2011")])
    only_d3 = b - a
    audit = {"method": "actual exact 8-column canonical tuple multiset, multiplicity preserved; no concatenation", "d2_rows": len(d2),
             "d3_rows": len(f), "d3_second_sheet_rows": sum(b.values()), "matched_rows": sum((a & b).values()),
             "d2_only_rows": sum((a - b).values()), "d3_second_sheet_only_rows": sum(only_d3.values()),
             "extra_rows_nonpersonal": [{"InvoiceNo": row[0], "StockCode": row[1], "Description": row[2], "quantity": row[3], "date": row[4], "unit_price": row[5], "country": row[7], "multiplicity": n} for row, n in only_d3.items()],
             "d2_duplicate_business_rows": len(d2) - len(a), "d3_duplicate_business_rows": int(duplicated.sum()),
             "sheet_counts": {k: int(v) for k, v in sheet_counts.items()}, "normalized_sha256": expected,
             "source_mapping_verified": "restored XLSX manifest hash (absPath removed) + normalized manifest hash + exact contiguous source sheet row mapping"}
    assert audit["matched_rows"] == 541909 and audit["d2_only_rows"] == 0 and audit["d3_second_sheet_only_rows"] == 1
    save_json(out / "overlap_audit.json", audit)
    sales = f[product].copy()
    sales["amount"] = sales.Quantity * sales.UnitPrice
    sales["month"] = sales.InvoiceDate.dt.strftime("%Y-%m")
    monthly = sales.groupby("month").agg(invoices=("InvoiceNo", "nunique"), units=("Quantity", "sum"), positive_sales_gbp=("amount", "sum")).reset_index()
    monthly["partial_month"] = monthly["month"].eq("2011-12")
    demand = demand_backtest(monthly)
    save_json(out / "demand_backtest.json", demand)
    known = sales[sales.CustomerID.ne("")]
    reference = sales.InvoiceDate.max().normalize() + pd.Timedelta(days=1)
    rfm = known.groupby("CustomerID").agg(last=("InvoiceDate", "max"), frequency=("InvoiceNo", "nunique"), monetary=("amount", "sum"))
    rfm["recency_days"] = (reference - rfm["last"].dt.normalize()).dt.days
    rfm_summary = {"customers": len(rfm), "reference_date": reference.strftime("%Y-%m-%d"),
                   "quantiles": {c: {str(q): float(v) for q, v in rfm[c].quantile([.25, .5, .75, .9]).items()} for c in ["recency_days", "frequency", "monetary"]},
                   "no_customer_rows_exported": True}
    # UK baskets (foreign gifts/wholesale) remain explicitly separate from baby catalog.
    uk = sales[sales.Country.eq("United Kingdom")]
    baskets = uk.groupby("InvoiceNo").agg(items=("StockCode", lambda v: sorted(set(v))), date=("InvoiceDate", "min")).sort_values(["date", "InvoiceNo"])
    cutoff = int(len(baskets) * .8)
    train, test = baskets.iloc[:cutoff], baskets.iloc[cutoff:]
    names = sales.sort_values("InvoiceDate").drop_duplicates("StockCode", keep="last").set_index("StockCode")["Description"].to_dict()
    item_counts = Counter(itertools.chain.from_iterable(train["items"]))
    top = {k for k, _ in item_counts.most_common(100)}
    pairs = Counter()
    for items in train["items"]:
        pairs.update(itertools.combinations(sorted(set(items) & top), 2))
    rules = []
    for (x, y), n in pairs.items():
        if n < 100:
            continue
        for left, right in [(x, y), (y, x)]:
            rules.append({"antecedent": left, "consequent": right, "antecedent_name": names[left], "consequent_name": names[right],
                          "joint_invoices": n, "antecedent_invoices": item_counts[left], "consequent_invoices": item_counts[right],
                          "total_invoices": len(train), "support": n / len(train), "confidence": n / item_counts[left],
                          "lift": n * len(train) / (item_counts[left] * item_counts[right]), "experimental": True,
                          "source": "D3 UCI Online Retail II / United Kingdom / chronological train invoices",
                          "source_url": "https://archive.ics.uci.edu/dataset/502/online+retail+ii",
                          "limitation": "영국 선물용품의 동시구매 실험이며 육아용품·현재 판매 상품 추천으로 연결하지 않음"})
    rules.sort(key=lambda v: (-v["lift"], -v["joint_invoices"], v["antecedent"], v["consequent"]))
    lookup = {}
    for r in rules:
        lookup.setdefault(r["antecedent"], []).append(r)
    hits_rules = hits_pop = evaluated = 0
    rr_rules = rr_pop = 0.0
    for invoice, row in test.iterrows():
        candidates = sorted(set(row["items"]) & top)
        if len(candidates) < 2:
            continue
        target = candidates[int(hashlib.sha256(invoice.encode()).hexdigest()[:8], 16) % len(candidates)]
        context = set(candidates) - {target}
        scores = {}
        for item in context:
            for r in lookup.get(item, []):
                if r["consequent"] not in context:
                    scores[r["consequent"]] = max(scores.get(r["consequent"], 0), r["confidence"])
        predicted = sorted(scores, key=lambda k: (-scores[k], k))[:5]
        popular = [k for k, _ in item_counts.most_common() if k not in context][:5]
        evaluated += 1
        for choice, is_rule in [(predicted, True), (popular, False)]:
            hit = target in choice
            rr = 1 / (choice.index(target) + 1) if hit else 0
            if is_rule:
                hits_rules += int(hit)
                rr_rules += rr
            else:
                hits_pop += int(hit)
                rr_pop += rr
    evaluation = {"task": "chronological held-out UK invoice basket completion, one deterministic held-out item, top100 train universe; remaining basket known",
                  "train_invoices": len(train), "test_invoices": len(test), "eligible_test_baskets": evaluated,
                  "train_last": train.date.max().isoformat(), "test_first": test.date.min().isoformat(),
                  "rules_recall_at_5": hits_rules / evaluated, "popularity_recall_at_5": hits_pop / evaluated,
                  "rules_mrr_at_5": rr_rules / evaluated, "popularity_mrr_at_5": rr_pop / evaluated,
                  "limitations": ["conditional on baskets with at least 2 training top100 items", "same historical customer may recur; no new-customer claim", "within-invoice completion is not next-order prediction", "no Korean childcare validation"]}
    save_json(out / "bundle_rules.json", rules)
    return {"audit": audit, "eligible_product_rows": len(sales), "invoices": int(sales.InvoiceNo.nunique()),
            "products": int(sales.StockCode.nunique()), "monthly": json.loads(monthly.to_json(orient="records")),
            "rfm": rfm_summary, "demand_backtest": demand, "bundle_rule_count": len(rules), "bundle_evaluation": evaluation,
            "limitations": ["historical UK gifts including wholesale, not Korean baby goods", "positive sales exclude cancellations and are not net revenue", "2011-12 is partial through December 9", "quantity is units, not group-buy people", "D2 is not appended to D3"]}


def prepare_splits(frame):
    x = frame[FEATURES].copy()
    x["Weekend"] = x["Weekend"].astype(int)
    x = x.astype(float)
    y = frame["Revenue"].astype(int).to_numpy()
    groups = pd.util.hash_pandas_object(x, index=False).to_numpy()
    outer = GroupShuffleSplit(n_splits=1, test_size=.2, random_state=SEED)
    development, test = next(outer.split(x, y, groups))
    inner = GroupShuffleSplit(n_splits=1, test_size=.25, random_state=SEED + 1)
    train_local, val_local = next(inner.split(x.iloc[development], y[development], groups[development]))
    return x, y, groups, {"train": development[train_local], "validation": development[val_local], "test": test}


def metrics(y, p):
    bins = []
    for low in np.arange(0, 1, .1):
        mask = (p >= low) & (p < low + .1 if low < .9 else p <= 1)
        if mask.any():
            bins.append({"lower": round(float(low), 1), "count": int(mask.sum()), "mean_probability": float(p[mask].mean()), "observed_rate": float(y[mask].mean())})
    return {"auroc": float(roc_auc_score(y, p)), "pr_auc_average_precision": float(average_precision_score(y, p)),
            "brier": float(brier_score_loss(y, p)), "log_loss": float(log_loss(y, p)),
            "ece_10bins": sum(v["count"] * abs(v["mean_probability"] - v["observed_rate"]) for v in bins) / len(y),
            "calibration_bins": bins}


def shoppers_analysis(path, out):
    f = pd.read_csv(path)
    x, y, groups, split = prepare_splits(f)
    assert all(not set(groups[split[a]]) & set(groups[split[b]]) for a, b in itertools.combinations(split, 2))
    models = {"dummy_prevalence": DummyClassifier(strategy="prior"),
              "logistic_regression": make_pipeline(SimpleImputer(strategy="median"), StandardScaler(), LogisticRegression(max_iter=1500, random_state=SEED)),
              "lightgbm": LGBMClassifier(n_estimators=160, learning_rate=.04, num_leaves=15, max_depth=6, min_child_samples=40,
                                          reg_lambda=2., n_jobs=2, random_state=SEED, deterministic=True, force_col_wise=True, verbosity=-1)}
    results = {}
    for name, model in models.items():
        model.fit(x.iloc[split["train"]], y[split["train"]])
        results[name] = {part: metrics(y[idx], model.predict_proba(x.iloc[idx])[:, 1]) for part, idx in split.items() if part != "train"}
    chosen = max(["logistic_regression", "lightgbm"], key=lambda n: results[n]["validation"]["pr_auc_average_precision"])
    joblib.dump(models[chosen], out / "intent_model.joblib")
    digest = lambda arr: hashlib.sha256(np.asarray(arr, dtype="int64").tobytes()).hexdigest()
    report = {"rows": len(f), "positive_labels": int(y.sum()), "label_prevalence": float(y.mean()), "features": FEATURES,
              "excluded_features": [c for c in f.columns if c not in FEATURES and c != "Revenue"],
              "seed": SEED, "split_method": "60/20/20 grouped by all selected feature values (not label); GroupShuffleSplit; validation selects model, test never selects",
              "duplicate_full_rows": int(f.duplicated().sum()), "duplicate_selected_feature_rows": int(pd.DataFrame(x).duplicated().sum()),
              "feature_group_overlap": 0, "splits": {k: {"rows": len(v), "positive_labels": int(y[v].sum()), "prevalence": float(y[v].mean()), "indices_sha256": digest(v)} for k, v in split.items()},
              "models": results, "selected_model": chosen, "selection_criterion": "maximum validation average precision among logistic and LightGBM",
              "model_sha256": sha256(out / "intent_model.joblib"), "calibration": "Uncalibrated probabilities evaluated with Brier/ECE and reliability bins on held-out test; no post-hoc calibration fit",
              "limitations": LIMITATIONS}
    save_json(out / "evaluation.json", report)
    save_json(out / "model_metadata.json", {k: report[k] for k in ["features", "selected_model", "model_sha256", "limitations", "seed"]})
    return report


def run(data_root, out):
    out.mkdir(parents=True, exist_ok=True)
    manifest = read_json(data_root / "notion_export/download_manifest.json")
    restored = read_json(data_root / "notion_export/restore_manifest.json")["datasets"]
    sources = []
    for dataset in manifest:
        for record in dataset["files"]:
            path = data_root / record["path"]
            actual = sha256(path)
            restored_record = next((v for v in restored if v["output_relative_path"] == record["path"]), None)
            expected = restored_record["output_sha256"] if restored_record else record["sha256"]
            assert actual == expected, f"Raw/restored source changed: {record['path']}"
            sources.append({"path": record["path"], "bytes": path.stat().st_size, "sha256": actual, "origin": dataset["source_url"],
                            "original_download_sha256": record["sha256"],
                            "copy_note": restored_record.get("note_ko", "original unchanged") if restored_record else "original unchanged",
                            "original_download_date": dataset["downloaded_at"], "current_check": "local file hash 2026-11-01; no new source download"})
    print("SOURCE_HASHES_PASS", flush=True)
    materials = material_inventory(data_root, out)
    print("MATERIALS_READ", materials, flush=True)
    d1 = consumer_analysis(data_root / "raw/01_ecommerce_behavior/Ecommerce_Consumer_Behavior_Analysis_Data.csv")
    print("D1_AGGREGATES", d1["rows"], flush=True)
    retail = retail_analysis(data_root, out)
    print("D2_D3_AUDIT", retail["audit"], flush=True)
    print("BASKET_EVALUATION", retail["bundle_evaluation"], flush=True)
    d4 = shoppers_analysis(data_root / "raw/04_uci_shoppers/online_shoppers_intention.csv", out)
    print("D4_SELECTED", d4["selected_model"], d4["models"][d4["selected_model"]]["test"], flush=True)
    summary = {"generated_at": datetime.now(timezone.utc).isoformat(), "experimental": True, "source_check_date": "2026-11-01",
               "sources": read_json(data_root / "notion_export/source_catalog.json"), "source_hashes": sources,
               "materials": materials, "d1": d1, "retail": retail, "intent": d4,
               "versions": {"python": platform.python_version(), "pandas": pd.__version__, "sklearn": sklearn.__version__},
               "privacy": "only aggregates/product association rules/models exported; no public customer record or app customer linkage",
               "scope": "Historical public-data experiments with provenance, separate from current baby-product catalog and production predictions"}
    save_json(out / "summary.json", summary)
    paths = [p for p in out.iterdir() if p.is_file() and p.name != "artifact_manifest.json"]
    save_json(out / "artifact_manifest.json", {"artifacts": [{"file": p.name, "sha256": sha256(p), "bytes": p.stat().st_size} for p in sorted(paths)]})
    print("PIPELINE_RUN_COMPLETE", flush=True)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--data-root", type=Path, default=HERE.parent / "_archive" / "legacy-project" / "data")
    parser.add_argument("--output", type=Path, default=HERE / "artifacts")
    args = parser.parse_args()
    run(args.data_root, args.output)
